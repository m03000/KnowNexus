"""笔记与个人文档目录元数据仓储。

Markdown 笔记和源文档仍保存在文件系统；本仓储只保存目录归属、来源关联与
拖拽位置。这样大文件不进入 SQLite，同时所有移动规则都有后端约束。
"""

from __future__ import annotations

import mimetypes
import re
import shutil
import uuid
import sqlite3
import os
import tempfile
from datetime import datetime
from pathlib import Path

from study_help_agent.infrastructure.persistence.sqlite import SqliteConnectionFactory


class LibraryCatalogRepository:
    EDITABLE_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".html", ".htm", ".css", ".csv", ".json", ".xml", ".yaml", ".yml"}
    EDITABLE_DOCUMENT_SUFFIXES = {".docx"}
    RESTRICTED_CREATE_FOLDERS = {"ai-note-root", "source-files"}

    def __init__(self, database: SqliteConnectionFactory, document_root: Path) -> None:
        self._database = database
        self._document_root = document_root.resolve()

    def ensure_note(self, filename: str) -> None:
        with self._database.connect() as connection:
            connection.execute(
                "INSERT OR IGNORE INTO note_library_entries(filename, folder_id) VALUES (?, 'ai-note-root')",
                (filename,),
            )

    def attach_source(self, filename: str, *, resource_id: str, source_name: str,
                      source_path: str, media_type: str) -> str | None:
        path = Path(source_path).resolve() if source_path else None
        if not path or not path.is_file() or media_type.startswith("video/"):
            return None
        document_id = resource_id or f"document_{uuid.uuid4().hex}"
        with self._database.connect() as connection:
            connection.execute(
                """INSERT OR IGNORE INTO library_documents
                   (document_id, folder_id, name, stored_path, media_type, size_bytes, resource_id, source_kind)
                   VALUES (?, 'source-files', ?, ?, ?, ?, ?, 'agent_source')""",
                (document_id, source_name or path.name, str(path), media_type, path.stat().st_size, resource_id),
            )
            connection.execute(
                "UPDATE note_library_entries SET source_document_id=? WHERE filename=?",
                (document_id, filename),
            )
        return document_id

    def tree(self, note_rows: list[dict]) -> dict:
        # Re-seed canonical folders for local databases created by older builds.
        with self._database.connect() as connection:
            connection.execute(
                "INSERT OR IGNORE INTO library_folders(folder_id,parent_id,library_type,name,is_system,sort_order) "
                "VALUES('ai-note-root',NULL,'ai_note','AI 笔记',1,0)"
            )
            connection.execute(
                "INSERT OR IGNORE INTO library_folders(folder_id,parent_id,library_type,name,is_system,sort_order) "
                "VALUES('personal-document-root',NULL,'personal_document','个人文档',1,1)"
            )
            connection.execute(
                "INSERT OR IGNORE INTO library_folders(folder_id,parent_id,library_type,name,is_system,sort_order) "
                "VALUES('source-files','personal-document-root','personal_document','源文件',1,0)"
            )
        for note in note_rows:
            self.ensure_note(note["filename"])
        with self._database.connect() as connection:
            folders = [dict(row) for row in connection.execute(
                "SELECT * FROM library_folders ORDER BY parent_id IS NOT NULL, sort_order, created_at"
            )]
            memberships = {
                row["filename"]: dict(row) for row in connection.execute(
                    "SELECT filename, folder_id, source_document_id FROM note_library_entries"
                )
            }
            documents = [dict(row) for row in connection.execute(
                "SELECT * FROM library_documents ORDER BY created_at DESC"
            )]
        notes = []
        for note in note_rows:
            membership = memberships.get(note["filename"], {})
            notes.append({**note, "item_type": "note", "item_id": note["filename"],
                          "folder_id": membership.get("folder_id", "ai-note-root"),
                          "source_document_id": membership.get("source_document_id")})
        for document in documents:
            path = Path(document["stored_path"])
            exists = path.is_file()
            stat = path.stat() if exists else None
            document.update(item_type="document", item_id=document["document_id"],
                            editable=exists and self._is_editable(path),
                            editor_format=self._editor_format(path) if exists else None,
                            missing=not exists,
                            size_bytes=stat.st_size if stat else document["size_bytes"],
                            modified_at=datetime.fromtimestamp(stat.st_mtime).isoformat() if stat else None)
        return {"folders": folders, "notes": notes, "documents": documents}

    def create_folder(self, parent_id: str, name: str) -> dict:
        cleaned = " ".join(name.split()).strip()
        if not cleaned or len(cleaned) > 80:
            raise ValueError("目录名称不能为空且不能超过 80 个字符")
        folder_id = f"folder_{uuid.uuid4().hex}"
        with self._database.connect() as connection:
            parent = connection.execute(
                "SELECT * FROM library_folders WHERE folder_id=?", (parent_id,)
            ).fetchone()
            if parent is None:
                raise ValueError("父目录不存在")
            if parent_id in self.RESTRICTED_CREATE_FOLDERS:
                raise ValueError("该系统目录不允许新建子目录")
            try:
                connection.execute(
                    "INSERT INTO library_folders(folder_id,parent_id,library_type,name) VALUES(?,?,?,?)",
                    (folder_id, parent_id, parent["library_type"], cleaned),
                )
            except sqlite3.IntegrityError as error:
                raise ValueError("同一目录下已存在同名目录") from error
        return {"folder_id": folder_id, "parent_id": parent_id,
                "library_type": parent["library_type"], "name": cleaned, "is_system": 0}

    def create_markdown_document(self, title: str, folder_id: str) -> dict:
        """在个人文档的自定义目录中创建一个可编辑的本地 Markdown 文件。"""
        cleaned = " ".join(title.split()).strip()
        if not cleaned:
            raise ValueError("文件名不能为空")
        if folder_id in self.RESTRICTED_CREATE_FOLDERS:
            raise ValueError("该系统目录不允许新建文件")
        safe_stem = Path(cleaned).stem.strip().rstrip(".")
        if not safe_stem or len(safe_stem) > 150 or re.search(r'[<>:"/\\|?*]', safe_stem):
            raise ValueError("文件名无效或过长")
        safe_name = f"{safe_stem}.md"
        with self._database.connect() as connection:
            target = connection.execute(
                "SELECT library_type FROM library_folders WHERE folder_id=?", (folder_id,)
            ).fetchone()
            if target is None or target["library_type"] != "personal_document":
                raise ValueError("Markdown 文件只能新建在个人文档目录中")
            duplicate = connection.execute(
                "SELECT 1 FROM library_documents WHERE folder_id=? AND lower(name)=lower(?)",
                (folder_id, safe_name),
            ).fetchone()
            if duplicate:
                raise ValueError("当前目录已存在同名文件")

        document_id = f"document_{uuid.uuid4().hex}"
        target_dir = self._document_root / document_id
        target_dir.mkdir(parents=True, exist_ok=False)
        stored = target_dir / safe_name
        stored.write_text(f"# {safe_stem}\n\n", encoding="utf-8")
        with self._database.connect() as connection:
            connection.execute(
                """INSERT INTO library_documents
                   (document_id,folder_id,name,stored_path,media_type,size_bytes,source_kind)
                   VALUES(?,?,?,?,?,?,'import')""",
                (document_id, folder_id, safe_name, str(stored), "text/markdown", stored.stat().st_size),
            )
        return {"document_id": document_id, "item_id": document_id, "item_type": "document",
                "folder_id": folder_id, "name": safe_name, "stored_path": str(stored),
                "media_type": "text/markdown", "size_bytes": stored.stat().st_size,
                "source_kind": "import", "editable": True, "editor_format": "text"}

    def delete_folder(self, folder_id: str) -> None:
        with self._database.connect() as connection:
            folder = connection.execute(
                "SELECT * FROM library_folders WHERE folder_id=?", (folder_id,)
            ).fetchone()
            if folder is None:
                raise FileNotFoundError(folder_id)
            if folder["is_system"]:
                raise ValueError("系统目录不能删除")
            occupied = connection.execute(
                """SELECT EXISTS(SELECT 1 FROM note_library_entries WHERE folder_id=?) OR
                          EXISTS(SELECT 1 FROM library_documents WHERE folder_id=?) OR
                          EXISTS(SELECT 1 FROM library_folders WHERE parent_id=?)""",
                (folder_id, folder_id, folder_id),
            ).fetchone()[0]
            if occupied:
                raise ValueError("目录非空，请先移动或删除其中内容")
            connection.execute("DELETE FROM library_folders WHERE folder_id=?", (folder_id,))

    def move_item(self, item_type: str, item_id: str, folder_id: str) -> None:
        table = "note_library_entries" if item_type == "note" else "library_documents"
        key = "filename" if item_type == "note" else "document_id"
        expected = "ai_note" if item_type == "note" else "personal_document"
        with self._database.connect() as connection:
            target = connection.execute(
                "SELECT library_type FROM library_folders WHERE folder_id=?", (folder_id,)
            ).fetchone()
            if target is None or target["library_type"] != expected:
                raise ValueError("内容只能在所属的一级目录内部移动")
            if item_type == "document" and folder_id == "personal-document-root":
                pass
            result = connection.execute(
                f"UPDATE {table} SET folder_id=? WHERE {key}=?", (folder_id, item_id)
            )
            if result.rowcount == 0:
                raise FileNotFoundError(item_id)

    def import_document(self, source: Path, name: str, media_type: str,
                        folder_id: str = "personal-document-root") -> dict:
        with self._database.connect() as connection:
            target = connection.execute(
                "SELECT library_type FROM library_folders WHERE folder_id=?", (folder_id,)
            ).fetchone()
            if target is None or target["library_type"] != "personal_document" or folder_id == "source-files":
                raise ValueError("导入文件只能放入个人文档目录（不能直接导入系统源文件目录）")
        document_id = f"document_{uuid.uuid4().hex}"
        safe_name = Path(name).name
        target_dir = self._document_root / document_id
        target_dir.mkdir(parents=True, exist_ok=False)
        stored = target_dir / safe_name
        shutil.copyfile(source, stored)
        detected = media_type or mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
        with self._database.connect() as connection:
            connection.execute(
                """INSERT INTO library_documents
                   (document_id,folder_id,name,stored_path,media_type,size_bytes,source_kind)
                   VALUES(?,?,?,?,?,?,'import')""",
                (document_id, folder_id, safe_name, str(stored), detected, stored.stat().st_size),
            )
        return {"document_id": document_id, "folder_id": folder_id, "name": safe_name,
                "media_type": detected, "size_bytes": stored.stat().st_size, "source_kind": "import"}

    def link_document(self, source: Path, name: str = "", media_type: str = "",
                      folder_id: str = "personal-document-root") -> dict:
        """只记录原文件绝对路径，不再向运行目录复制第二份文件。"""
        path = source.expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(str(source))
        with self._database.connect() as connection:
            target = connection.execute(
                "SELECT library_type FROM library_folders WHERE folder_id=?", (folder_id,)
            ).fetchone()
            if target is None or target["library_type"] != "personal_document" or folder_id == "source-files":
                raise ValueError("关联文件只能放入个人文档目录")
            existing = connection.execute(
                "SELECT * FROM library_documents WHERE stored_path=?", (str(path),)
            ).fetchone()
            if existing is not None:
                connection.execute(
                    "UPDATE library_documents SET folder_id=? WHERE document_id=?",
                    (folder_id, existing["document_id"]),
                )
                return {**dict(existing), "folder_id": folder_id, "already_linked": True}

        document_id = f"document_{uuid.uuid4().hex}"
        safe_name = Path(name).name if name else path.name
        detected = media_type or mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
        with self._database.connect() as connection:
            connection.execute(
                """INSERT INTO library_documents
                   (document_id,folder_id,name,stored_path,media_type,size_bytes,source_kind)
                   VALUES(?,?,?,?,?,?,'linked')""",
                (document_id, folder_id, safe_name, str(path), detected, path.stat().st_size),
            )
        return {"document_id": document_id, "item_id": document_id, "item_type": "document",
                "folder_id": folder_id, "name": safe_name, "stored_path": str(path),
                "media_type": detected, "size_bytes": path.stat().st_size, "source_kind": "linked"}

    def document(self, document_id: str) -> dict:
        with self._database.connect() as connection:
            row = connection.execute(
                "SELECT * FROM library_documents WHERE document_id=?", (document_id,)
            ).fetchone()
        if row is None or not Path(row["stored_path"]).is_file():
            raise FileNotFoundError(document_id)
        result = dict(row)
        path = Path(result["stored_path"])
        result["editable"] = self._is_editable(path)
        result["editor_format"] = self._editor_format(path)
        return result

    @classmethod
    def _is_editable(cls, path: Path) -> bool:
        suffix = path.suffix.lower()
        return suffix in cls.EDITABLE_TEXT_SUFFIXES or suffix in cls.EDITABLE_DOCUMENT_SUFFIXES

    @classmethod
    def _editor_format(cls, path: Path) -> str | None:
        if path.suffix.lower() == ".docx":
            return "docx-markdown"
        return "text" if path.suffix.lower() in cls.EDITABLE_TEXT_SUFFIXES else None

    def read_document_text(self, document_id: str) -> str:
        """返回编辑器中的可往返文本。

        DOCX 使用 Markdown 作为中间格式，Heading 与列表样式会转成标记，
        让前端能实时生成目录。
        """
        document = self.document(document_id)
        if not document["editable"]:
            raise ValueError("该文件格式只读，PDF、旧版 Office 文档和图片暂不允许编辑")
        path = Path(document["stored_path"])
        if path.suffix.lower() == ".docx":
            return self._read_docx_as_markdown(path)
        return path.read_text(encoding="utf-8", errors="replace")

    def update_document_text(self, document_id: str, content: str) -> dict:
        """更新可往返保存的文本或 DOCX 资源，并同步文件大小。"""
        document = self.document(document_id)
        if not document["editable"]:
            raise ValueError("该文件格式只读，PDF、旧版 Office 文档和图片暂不允许编辑")
        path = Path(document["stored_path"])
        if path.suffix.lower() == ".docx":
            self._update_docx_text_safely(
                path, content, create_backup=document.get("source_kind") != "linked"
            )
            self._update_document_size(document_id, path)
            return self.document(document_id)
        fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="") as stream:
                stream.write(content)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            Path(temporary).unlink(missing_ok=True)
        self._update_document_size(document_id, path)
        return self.document(document_id)

    def _update_document_size(self, document_id: str, path: Path) -> None:
        with self._database.connect() as connection:
            connection.execute(
                "UPDATE library_documents SET size_bytes=? WHERE document_id=?",
                (path.stat().st_size, document_id),
            )

    @staticmethod
    def _read_docx_as_markdown(path: Path) -> str:
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("当前环境未安装 python-docx，无法编辑 DOCX") from error

        document = Document(path)
        blocks: list[str] = []
        numbering_counters: dict[tuple[int, int], int] = {}
        numbering_starts: dict[tuple[int, int], int] = {}
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name if paragraph.style else "").lower()
            heading = re.search(r"heading\s*([1-6])", style_name)
            if heading:
                blocks.append(f"{'#' * int(heading.group(1))} {text}")
            elif "list bullet" in style_name:
                blocks.append(f"- {text}")
            elif "list number" in style_name:
                key, start = LibraryCatalogRepository._docx_numbering_identity(
                    document, paragraph, numbering_starts
                )
                current = numbering_counters.get(key, start - 1) + 1
                numbering_counters[key] = current
                blocks.append(f"{current}. {text}")
            else:
                blocks.append(text)
            blocks.append("")
        return "\n".join(blocks).strip() + "\n"

    @staticmethod
    def _docx_numbering_identity(document, paragraph, starts: dict[tuple[int, int], int]) -> tuple[tuple[int, int], int]:
        """读取 Word 编号定义，避免把每个 List Number 段落都转换成 1。

        Word 不会把显示编号写进 paragraph.text，编号由 numId + ilvl
        以及 numbering.xml 中的 start 共同决定。
        """
        from docx.oxml.ns import qn

        direct = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
        styled = paragraph.style.element.pPr.numPr if paragraph.style and paragraph.style.element.pPr is not None else None
        num_pr = direct or styled
        num_id = int(num_pr.numId.val) if num_pr is not None and num_pr.numId is not None else -1
        level = int(num_pr.ilvl.val) if num_pr is not None and num_pr.ilvl is not None else 0
        key = (num_id, level)
        if key in starts or num_id < 0:
            return key, starts.get(key, 1)

        start = 1
        numbering = document.part.numbering_part.element
        number = next(
            (node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == str(num_id)),
            None,
        )
        if number is not None:
            override = next(
                (node for node in number.findall(qn("w:lvlOverride")) if node.get(qn("w:ilvl")) == str(level)),
                None,
            )
            start_override = override.find(qn("w:startOverride")) if override is not None else None
            if start_override is not None:
                start = int(start_override.get(qn("w:val"), "1"))
            abstract_id_node = number.find(qn("w:abstractNumId"))
            abstract_id = abstract_id_node.get(qn("w:val")) if abstract_id_node is not None else None
            abstract = next(
                (node for node in numbering.findall(qn("w:abstractNum")) if node.get(qn("w:abstractNumId")) == abstract_id),
                None,
            )
            level_node = next(
                (node for node in abstract.findall(qn("w:lvl")) if node.get(qn("w:ilvl")) == str(level)),
                None,
            ) if abstract is not None else None
            level_start = level_node.find(qn("w:start")) if level_node is not None else None
            if start_override is None and level_start is not None:
                start = int(level_start.get(qn("w:val"), "1"))
        starts[key] = start
        return key, start

    @staticmethod
    def _update_docx_text_safely(path: Path, content: str, *, create_backup: bool = True) -> None:
        """仅更新 DOCX 正文段落文字，不重建文档。

        图片、表格、分节、页眉页脚保持在原文档中；编辑器新增的段落追加到
        正文末尾。第一次保存还会留下 .original.docx 恢复副本。
        """
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("当前环境未安装 python-docx，无法保存 DOCX") from error

        backup = path.with_name(f"{path.stem}.original{path.suffix}")
        if create_backup and not backup.exists():
            shutil.copy2(path, backup)

        document = Document(path)
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        edited_blocks = [line.strip() for line in content.splitlines() if line.strip()]

        for index, raw in enumerate(edited_blocks):
            heading = re.match(r"^(#{1,6})\s+(.+)$", raw)
            bullet = re.match(r"^[-*+]\s+(.+)$", raw)
            numbered = re.match(r"^\d+[.)]\s+(.+)$", raw)
            text = heading.group(2) if heading else bullet.group(1) if bullet else numbered.group(1) if numbered else raw
            paragraph = paragraphs[index] if index < len(paragraphs) else document.add_paragraph()
            if heading:
                paragraph.style = f"Heading {len(heading.group(1))}"
            elif bullet:
                paragraph.style = "List Bullet"
            elif numbered:
                paragraph.style = "List Number"
            elif paragraph.style and (
                paragraph.style.name.lower().startswith("heading")
                or "list " in paragraph.style.name.lower()
            ):
                paragraph.style = "Normal"
            if paragraph.runs:
                paragraph.runs[0].text = text
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(text)

        # 编辑器删除的文字仅清空对应段落，不删除 XML 节点，避免误伤图片和布局。
        for paragraph in paragraphs[len(edited_blocks):]:
            for run in paragraph.runs:
                run.text = ""

        fd, temporary = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".docx", dir=path.parent)
        os.close(fd)
        try:
            document.save(temporary)
            with open(temporary, "rb+") as stream:
                os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            Path(temporary).unlink(missing_ok=True)

    def delete_document(self, document_id: str) -> None:
        """删除目录记录和受控保存的文件；Agent 生成笔记所依赖的源文件禁止误删。"""
        with self._database.connect() as connection:
            row = connection.execute(
                "SELECT * FROM library_documents WHERE document_id=?", (document_id,)
            ).fetchone()
            if row is None:
                raise FileNotFoundError(document_id)
            if row["source_kind"] == "agent_source":
                raise ValueError("笔记关联的系统源文件不能从个人文档接口直接删除")
            connection.execute(
                "DELETE FROM library_documents WHERE document_id=?", (document_id,)
            )
        path = Path(row["stored_path"])
        # linked 条目只解除关联，永远不删除用户的原文件。
        if row["source_kind"] == "linked":
            return
        path.unlink(missing_ok=True)
        if path.suffix.lower() == ".docx":
            path.with_name(f"{path.stem}.original{path.suffix}").unlink(missing_ok=True)
        parent = path.parent.resolve()
        if parent != self._document_root and self._document_root in parent.parents:
            try:
                parent.rmdir()
            except OSError:
                pass

    def dashboard(self, note_count: int) -> dict:
        with self._database.connect() as connection:
            document_count = connection.execute("SELECT COUNT(*) FROM library_documents").fetchone()[0]
            source_count = connection.execute(
                "SELECT COUNT(*) FROM library_documents WHERE source_kind='agent_source'"
            ).fetchone()[0]
            imported_count = document_count - source_count
            note_folders = connection.execute(
                "SELECT COUNT(*) FROM library_folders WHERE library_type='ai_note' AND parent_id IS NOT NULL"
            ).fetchone()[0]
            document_folders = connection.execute(
                """SELECT COUNT(*) FROM library_folders
                   WHERE library_type='personal_document' AND parent_id IS NOT NULL AND is_system=0"""
            ).fetchone()[0]
        return {"ai_notes": {"count": note_count, "folders": note_folders},
                "personal_documents": {"count": document_count, "imported": imported_count,
                                       "source_files": source_count, "folders": document_folders}}

    def remove_note(self, filename: str) -> None:
        with self._database.connect() as connection:
            connection.execute("DELETE FROM note_library_entries WHERE filename=?", (filename,))
