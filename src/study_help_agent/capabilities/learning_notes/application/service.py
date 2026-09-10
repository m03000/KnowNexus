import re
import logging
import os
import tempfile
from pathlib import Path

from study_help_agent.capabilities.learning_notes.application.point_service import (
    KnowledgePointService,
)
from study_help_agent.capabilities.learning_notes.application.ports import NoteRepository
from study_help_agent.capabilities.learning_notes.domain.models import (
    BuiltLearningNote,
    GeneratedNote,
    NoteSegment,
    SavedNote,
)
from study_help_agent.core.exceptions import LearningNoteNotFoundError
from study_help_agent.capabilities.learning_notes.infrastructure.library_catalog_repository import (
    LibraryCatalogRepository,
)
from study_help_agent.capabilities.learning_notes.application.document_ingestion_service import (
    ImportedDocumentIngestionService,
)
from study_help_agent.capabilities.knowledge_ingestion.application.deletion_service import (
    KnowledgeDeletionService,
)
from study_help_agent.capabilities.llm_wiki.application import WikiSourceService

logger = logging.getLogger(__name__)


class LearningNoteService:
    def __init__(self, *, repository: NoteRepository,
                 point_service: KnowledgePointService | None = None,
                 catalog: LibraryCatalogRepository | None = None,
                 document_ingestion: ImportedDocumentIngestionService | None = None,
                 knowledge_deletion: KnowledgeDeletionService | None = None,
                 wiki_sources: WikiSourceService | None = None,
                 export_directory: Path | None = None,
                 ) -> None:
        self._repository = repository
        self._point_service = point_service
        self._catalog = catalog
        self._document_ingestion = document_ingestion
        self._knowledge_deletion = knowledge_deletion
        self._wiki_sources = wiki_sources
        self._export_directory = export_directory

    def list_notes(self) -> dict:
        """业务层实现的查询所有笔记功能"""
        return {"notes": self._repository.list_notes()}

    def library_tree(self) -> dict:
        return self._catalog.tree(self._repository.list_notes()) if self._catalog else {
            "folders": [], "notes": self._repository.list_notes(), "documents": []
        }

    def create_folder(self, parent_id: str, name: str) -> dict:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        return self._catalog.create_folder(parent_id, name)

    def delete_folder(self, folder_id: str) -> None:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        self._catalog.delete_folder(folder_id)

    def create_markdown_document(self, title: str, folder_id: str) -> dict:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        return self._catalog.create_markdown_document(title, folder_id)

    def move_library_item(self, item_type: str, item_id: str, folder_id: str) -> None:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        self._catalog.move_item(item_type, item_id, folder_id)

    def import_document(self, source_path, name: str, media_type: str, folder_id: str) -> dict:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        document = self._catalog.import_document(source_path, name, media_type, folder_id)
        if self._document_ingestion is None:
            return document
        try:
            ingestion = self._document_ingestion.ingest(self._catalog.document(document["document_id"]))
        except Exception:
            # 文档管理与知识索引是两个生命周期：提取失败不应丢掉用户已导入的原文件。
            logger.exception("个人文档自动入库失败：%s", document["document_id"])
            return {**document, "rag_ingestion": {"status": "failed"}}
        return {**document, "rag_ingestion": {"status": "queued", **ingestion}}

    def link_document(self, source_path: str, name: str, media_type: str, folder_id: str) -> dict:
        """关联原文件并增量入库，不制造第二份文件。"""
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        document = self._catalog.link_document(Path(source_path), name, media_type, folder_id)
        if self._document_ingestion is None:
            return document
        try:
            ingestion = self._document_ingestion.ingest(self._catalog.document(document["document_id"]))
        except Exception:
            logger.exception("关联文档自动入库失败：%s", document["document_id"])
            return {**document, "rag_ingestion": {"status": "failed"}}
        return {**document, "rag_ingestion": {"status": "queued", **ingestion}}

    def reindex_document(self, document_id: str) -> dict:
        """重新提交同一稳定文档；内容未变化时由仓储幂等跳过。"""
        if not self._catalog or not self._document_ingestion:
            raise RuntimeError("文档入库服务未配置")
        result = self._document_ingestion.ingest(self._catalog.document(document_id))
        return {"document_id": document_id, "rag_ingestion": result}

    def analyze_document_knowledge(self, document_id: str) -> dict:
        """在用户明确同意后，从个人文档提取知识点并关联到笔记图谱。"""
        if not self._catalog or not self._document_ingestion or not self._point_service:
            raise RuntimeError("文档知识点分析服务未配置")
        document = self._catalog.document(document_id)
        source_key = f"document:{document_id}"
        existing = self._point_service.points_for_source(source_key)
        if existing:
            return {
                "document_id": document_id,
                "status": "already_analyzed",
                "knowledge_point_count": len(existing),
            }
        _, _, content = self._document_ingestion.extract(document)
        segments = self._split_markdown_segments(content)
        result = self._point_service.extract_and_link(source_key, segments)
        return {
            "document_id": document_id,
            "status": "completed",
            "knowledge_point_count": len(result.links),
            "new_knowledge_point_count": len(result.knowledge_points),
        }

    def delete_document(self, document_id: str) -> dict:
        """删除个人文档，并可靠排队清理其关键词和向量索引。"""
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        document = self._catalog.document(document_id)
        if document.get("source_kind") == "agent_source":
            raise ValueError("笔记关联的系统源文件不能从个人文档接口直接删除")
        queued = self._knowledge_deletion.delete_document(document_id) if self._knowledge_deletion else 0
        if self._point_service is not None:
            self._point_service.delete_links_for_note(f"document:{document_id}")
        if self._wiki_sources is not None:
            try:
                self._wiki_sources.delete(f"document:{document_id}")
            except Exception:
                logger.exception("Wiki 文档来源删除排队失败：%s", document_id)
        self._catalog.delete_document(document_id)
        action = "已解除关联" if document.get("source_kind") == "linked" else "已删除"
        return {
            "message": f"{document['name']} {action}",
            "knowledge_assets_queued_for_deletion": queued,
        }

    def document(self, document_id: str) -> dict:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        return self._catalog.document(document_id)

    def document_content(self, document_id: str) -> dict:
        document = self.document(document_id)
        if not document.get("editable"):
            return {**document, "content": None}
        return {**document, "content": self._catalog.read_document_text(document_id)}

    def update_document_content(self, document_id: str, content: str) -> dict:
        if not self._catalog:
            raise RuntimeError("目录服务未配置")
        document = self._catalog.update_document_text(document_id, content)
        ingestion = self._document_ingestion.ingest(document) if self._document_ingestion else None
        return {"message": "已保存", "document": document, "rag_ingestion": ingestion}

    def dashboard(self) -> dict:
        return self._catalog.dashboard(len(self._repository.list_notes())) if self._catalog else {}

    def publish(self, note: BuiltLearningNote) -> SavedNote:
        """业务层实现的保存功能：将已完成的生成结果默认保存为前端可直接读取的 Markdown 笔记。"""

        body = note.content.strip()
        if not body.startswith("# "):
            body = f"# {note.title}\n\n{body}"
        if note.introduction.strip():
            lines = body.splitlines()
            insert_at = 1 if lines and lines[0].startswith("# ") else 0
            lines[insert_at:insert_at] = ["", f"> {note.introduction.strip()}", ""]
            body = "\n".join(lines)
        source_uri = str(note.metadata.get("source_uri") or "").strip()
        resource_id = str(note.metadata.get("resource_id") or "").strip()
        media_type = str(note.metadata.get("media_type") or "").strip()
        source_name = str(note.metadata.get("source_name") or "").strip()
        source_lines: list[str] = []
        if source_uri.startswith(("http://", "https://")):
            source_lines.append(f"> 原始来源：[打开外部资源]({source_uri})")
        if resource_id and not media_type.startswith("video/"):
            source_lines.append(
                f"> [源文件](/api/library/notes/documents/{resource_id})"
            )
        if source_lines:
            lines = body.splitlines()
            insert_at = 1 if lines and lines[0].startswith("# ") else 0
            lines[insert_at:insert_at] = ["", *source_lines, ""]
            body = "\n".join(lines)
        segments = self._split_markdown_segments(body)
        saved = self._repository.save(
            GeneratedNote(
                title=note.title,
                content=body,
                topics=list(note.topics),
                segments = segments,
            )
        )
        if self._catalog is not None:
            self._catalog.ensure_note(saved.filename)
            self._catalog.attach_source(
                saved.filename,
                resource_id=resource_id,
                source_name=source_name,
                source_path=str(note.metadata.get("managed_source_path") or ""),
                media_type=media_type,
            )
        # 先拿到笔记名称，再进行知识点提取
        knowledge_points_refreshed = False
        if self._point_service is not None:
            try:
                self._point_service.extract_and_link(saved.filename, segments)
                knowledge_points_refreshed = True
            except Exception as error:  # 知识点提取失败不阻塞笔记保存
                logger.warning("知识点提取失败：%s（笔记已保存为 %s）", error, saved.filename)
        if self._wiki_sources is not None:
            self._sync_wiki_note(
                saved.filename,
                note.title,
                body,
                metadata={
                    "topics": list(note.topics),
                    "source_uri": source_uri,
                    "resource_id": resource_id,
                    "quality_status": note.quality_status,
                    "knowledge_points_refreshed": knowledge_points_refreshed,
                },
            )
        return saved

    def read(self, filename: str) -> dict:
        """读取完整笔记及其分段索引，供图谱跳转后精确定位。"""
        try:
            content = self._repository.read(filename)
        except FileNotFoundError as error:
            raise LearningNoteNotFoundError(f"学习笔记不存在：{filename}") from error
        segments = self._split_markdown_segments(content)
        return {
            "filename": filename,
            "content": content,
            "segments": [
                {
                    "segment_index": index,
                    "title": segment.title,
                    "content": segment.content,
                }
                for index, segment in enumerate(segments)
            ],
        }

    def create_manual_note(self, title: str, folder_id: str = "ai-note-root") -> dict:
        title = " ".join(title.split()).strip()
        if not title:
            raise ValueError("笔记标题不能为空")
        saved = self._repository.save(GeneratedNote(title=title, content=f"# {title}\n\n", topics=[], segments=[]))
        if self._catalog:
            self._catalog.ensure_note(saved.filename)
            if folder_id != "ai-note-root":
                self._catalog.move_item("note", saved.filename, folder_id)
        if self._wiki_sources is not None:
            self._sync_wiki_note(saved.filename, title, f"# {title}\n\n")
        return {"filename": saved.filename, "item_id": saved.filename, "item_type": "note", "folder_id": folder_id, "title": title}

    def update(self, filename: str, content: str, *, refresh_knowledge: bool = False) -> dict:
        self._repository.update(filename, content)

        segments = self._split_markdown_segments(content)

        # 自动保存不能每次按键都调用 LLM；需要时由显式“刷新知识点”动作触发。
        if refresh_knowledge and self._point_service is not None:
            result = self._point_service.extract_and_link(filename, segments)
            self._point_service.retain_extracted_links(filename, result.links)

        if self._wiki_sources is not None:
            title_match = re.search(r"^#\s+(.+)$", content, flags=re.MULTILINE)
            title = title_match.group(1).strip() if title_match else Path(filename).stem
            self._sync_wiki_note(
                filename,
                title,
                content,
                metadata={
                    "wiki_build_delay_seconds": 10,
                    "knowledge_points_refreshed": refresh_knowledge,
                },
            )

        return {
            "message": "已保存",
            "segments_count": len(segments),
            "knowledge_refresh_pending": not refresh_knowledge,
        }

    def export_note(self, filename: str, export_format: str) -> dict:
        """将系统生成的 Markdown 笔记导出到固定目录。"""
        if self._export_directory is None:
            raise RuntimeError("笔记导出目录未配置")
        content = self._repository.read(filename)
        export_format = export_format.lower()
        if export_format not in {"md", "docx"}:
            raise ValueError("暂时只支持导出 MD 和 DOCX")
        self._export_directory.mkdir(parents=True, exist_ok=True)
        stem = Path(filename).stem
        target = (self._export_directory / f"{stem}.{export_format}").resolve()
        if target.parent != self._export_directory.resolve():
            raise ValueError("非法的导出文件名")
        if export_format == "md":
            self._atomic_export_text(target, content)
        else:
            self._export_markdown_docx(target, content)
        return {"message": "导出完成", "format": export_format, "path": str(target)}

    @staticmethod
    def _atomic_export_text(target: Path, content: str) -> None:
        fd, temporary = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp", dir=target.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="") as stream:
                stream.write(content)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, target)
        finally:
            Path(temporary).unlink(missing_ok=True)

    @staticmethod
    def _export_markdown_docx(target: Path, content: str) -> None:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        # DOCX 导出采用 compact_reference_guide：适合信息密度较高的个人笔记。
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(0.492)
        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        heading_tokens = {
            1: (16, "2E74B5", 18, 10),
            2: (13, "2E74B5", 14, 7),
            3: (12, "1F4D78", 10, 5),
            4: (11, "1F4D78", 8, 4),
            5: (11, "1F4D78", 7, 3),
            6: (11, "1F4D78", 6, 3),
        }
        for level in range(1, 7):
            style = document.styles[f"Heading {level}"]
            size, color, before, after = heading_tokens[level]
            style.font.name = "Calibri"
            style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            style.font.color.rgb = RGBColor.from_string(color)
            style.font.size = Pt(size)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
        for style_name in ("List Bullet", "List Number"):
            list_style = document.styles[style_name]
            list_style.paragraph_format.left_indent = Inches(0.375)
            list_style.paragraph_format.first_line_indent = Inches(-0.188)
            list_style.paragraph_format.space_after = Pt(4)
            list_style.paragraph_format.line_spacing = 1.25

        for raw_line in content.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", line)
            bullet = re.match(r"^[-*+]\s+(.+)$", line)
            numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
            quote = re.match(r"^>\s*(.+)$", line)
            if heading:
                paragraph = document.add_paragraph(style=f"Heading {len(heading.group(1))}")
                text = heading.group(2)
            elif bullet:
                paragraph = document.add_paragraph(style="List Bullet")
                text = bullet.group(1)
            elif numbered:
                paragraph = document.add_paragraph(style="List Number")
                text = numbered.group(1)
            elif quote:
                paragraph = document.add_paragraph(style="Quote")
                text = quote.group(1)
            else:
                paragraph = document.add_paragraph()
                text = line
            LearningNoteService._append_markdown_runs(paragraph, text)

        fd, temporary = tempfile.mkstemp(prefix=f".{target.stem}.", suffix=".docx", dir=target.parent)
        os.close(fd)
        try:
            document.save(temporary)
            os.replace(temporary, target)
        finally:
            Path(temporary).unlink(missing_ok=True)

    @staticmethod
    def _append_markdown_runs(paragraph, text: str) -> None:
        token_pattern = re.compile(r"(\*\*.+?\*\*|__.+?__|`.+?`|(?<!\*)\*[^*]+?\*(?!\*))")
        cursor = 0
        for match in token_pattern.finditer(text):
            if match.start() > cursor:
                paragraph.add_run(text[cursor:match.start()])
            token = match.group(0)
            run = paragraph.add_run(token[2:-2] if token.startswith(("**", "__")) else token[1:-1])
            if token.startswith(("**", "__")):
                run.bold = True
            elif token.startswith("*"):
                run.italic = True
            else:
                run.font.name = "Cascadia Mono"
            cursor = match.end()
        if cursor < len(text):
            paragraph.add_run(text[cursor:])

    def delete(self, filename: str) -> dict:
        """业务层实现的删除功能：删除指定的笔记文件以及对应的RAG数据"""
        if not self._repository.delete(filename):
            raise LearningNoteNotFoundError(f"学习笔记不存在：{filename}")
        if self._point_service is not None:
            self._point_service.delete_links_for_note(filename)
        if self._catalog is not None:
            self._catalog.remove_note(filename)
        if self._wiki_sources is not None:
            try:
                self._wiki_sources.delete(f"note:{filename}")
            except Exception:
                logger.exception("Wiki 笔记来源删除排队失败：%s", filename)
        return {"message": f"{filename} 已删除"}

    def _sync_wiki_note(
        self,
        filename: str,
        title: str,
        content: str,
        *,
        metadata: dict | None = None,
    ) -> None:
        """Wiki 故障不得破坏笔记保存主流程；失败会保留日志供后续重建。"""
        if self._wiki_sources is None:
            return
        try:
            self._wiki_sources.upsert_note(
                filename,
                title,
                content,
                metadata=metadata,
            )
        except Exception:
            logger.exception("Wiki 笔记来源同步失败：%s", filename)

    def graph_data(self) -> dict:
        """返回笔记图谱（知识点-笔记），供统一图谱查询入口使用。"""
        graph = (
            self._point_service.get_graph_data()
            if self._point_service is not None
            else {"nodes": [], "edges": []}
        )
        nodes = list(graph.get("nodes", []))
        nodes_by_id = {str(node.get("id")): node for node in nodes}
        if self._catalog is not None:
            documents = {
                str(item["document_id"]): item
                for item in self._catalog.tree(self._repository.list_notes()).get("documents", [])
            }
            for node in nodes:
                filename = str(node.get("payload", {}).get("filename") or "")
                if not filename.startswith("document:"):
                    continue
                document_id = filename.removeprefix("document:")
                document = documents.get(document_id)
                if document is None:
                    continue
                title = str(document.get("name") or "个人文档")
                node.update({
                    "label": title,
                    "node_type": "document",
                    "summary": f"个人文档：{title}",
                })
                node.setdefault("payload", {}).update({
                    "document_id": document_id,
                    "source_kind": document.get("source_kind"),
                })
        for note in self._repository.list_notes():
            node_id = f"note:{note['filename']}"
            if node_id in nodes_by_id:
                node = nodes_by_id[node_id]
                node["label"] = note["title"]
                node["summary"] = f"学习笔记：{note['title']}"
                node.setdefault("payload", {}).update({
                    "filename": note["filename"],
                    "created_at": note.get("modified"),
                })
                continue
            node = {
                "id": node_id,
                "label": note["title"],
                "node_type": "note",
                "domain": "note",
                "summary": f"学习笔记：{note['title']}",
                "payload": {
                    "filename": note["filename"],
                    "created_at": note.get("modified"),
                },
            }
            nodes.append(node)
            nodes_by_id[node_id] = node
        return {"nodes": nodes, "edges": list(graph.get("edges", []))}

    @staticmethod
    def _split_markdown_segments(body: str) -> list[NoteSegment]:
        """把最终 Markdown 按标题(#/##/###)切成段落，供知识点精确定位。

        返回的列表下标即 NotePointLink.segment_index 的锚点；
        每段 title 是标题文本，方便前端定位到标题锚点。
        """
        segments: list[NoteSegment] = []
        current_title = ""
        current_lines: list[str] = []
        for line in body.splitlines():
            if re.match(r"^#{1,3} ", line):
                if current_lines:
                    segments.append(
                        NoteSegment(
                            category="markdown",
                            title=current_title,
                            content="\n".join(current_lines),
                            sort_order=len(segments),
                        )
                    )
                current_title = line.lstrip("# ").strip()
                current_lines = [line]
            else:
                current_lines.append(line)
        if current_lines:
            segments.append(
                NoteSegment(
                    category="markdown",
                    title=current_title,
                    content="\n".join(current_lines),
                    sort_order=len(segments),
                )
            )
        return segments

    def read_with_segments(self, filename: str) -> dict:
        return self.read(filename)
